"""Rebuild the report with the bundled workspace Python and python-docx.

Reads only the dated research inputs. No network calls or product mutations.
Run this file with the Python path from load_workspace_dependencies, then run
the documents skill render_docx.py and inspect every page before distribution.
"""
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
report = json.loads((ROOT / 'report.json').read_text())
sources = json.loads((ROOT / 'sources.json').read_text())
by_id = {s['id']: s for s in sources}
runs = {p.stem: json.loads(p.read_text()) for p in sorted((ROOT / 'data').glob('*.json'))}
metrics = {n: {r['keyword']: r for r in d['results']} for n, d in runs.items()}
doc = Document()
for el in list(doc.styles.element.iter(qn('w:pBdr'))):
    el.getparent().remove(el)
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.2677), Inches(11.6929)
sec.top_margin = sec.bottom_margin = Inches(.72)
sec.left_margin = sec.right_margin = Inches(.76)
sec.header_distance = sec.footer_distance = Inches(.32)
WIDTH = 8.2677 - 1.52
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.widow_control = True
normal = doc.styles['Normal']
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12
for name, size, before, after in [('Title', 25, 0, 14), ('Subtitle', 12, 0, 10), ('Heading 1', 16, 17, 9), ('Heading 2', 12, 12, 6)]:
    st = doc.styles[name]
    st.font.size = Pt(size)
    st.font.bold = name != 'Subtitle'
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
doc.core_properties.title = report['title']
doc.core_properties.subject = 'Marktbedarf und Capability Match für LinkedIn MCP Bridge'
doc.core_properties.author = 'Codex für Hagen'
doc.core_properties.keywords = 'LinkedIn, MCP, Marktstudie, Claude, ChatGPT, Google Ads'

def hyperlink(p, text, url):
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    col = OxmlElement('w:color'); col.set(qn('w:val'), '244A64'); props.append(col)
    run.append(props)
    txt = OxmlElement('w:t'); txt.text = text; run.append(txt)
    link.append(run); p._p.append(link)

def rich(p, text):
    text = text.replace('notarialisierte', 'notarisierte')
    for part in re.split(r'(\[S\d{2}(?:, S\d{2})*\])', text):
        if re.fullmatch(r'\[S\d{2}(?:, S\d{2})*\]', part):
            p.add_run('[')
            for i, sid in enumerate(part[1:-1].split(', ')):
                if i: p.add_run(', ')
                hyperlink(p, sid, by_id[sid]['url'])
            p.add_run(']')
        else:
            p.add_run(part)
    return p

def p(text, style=None):
    return rich(doc.add_paragraph(style=style), text)

def table(headers, rows, widths=None, size=9):
    widths = widths or [1 / len(headers)] * len(headers)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, fraction in zip(t.columns, widths): col.width = Inches(WIDTH * fraction)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + name)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]: el.set(qn('w:'+k),v)
        borders.append(el)
    props.append(borders)
    allrows = [headers] + rows
    for ri, values in enumerate(allrows):
        row = t.rows[0] if ri == 0 else t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if ri == 0: trpr.append(OxmlElement('w:tblHeader'))
        for ci, (cell, text) in enumerate(zip(row.cells, values)):
            cell.width = Inches(WIDTH * widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '243E50' if ri == 0 else ('F3F6F8' if ri % 2 == 0 else 'FFFFFF')); tcpr.append(shade)
            margins = OxmlElement('w:tcMar')
            for side, val in [('top','90'),('bottom','90'),('left','100'),('right','100')]:
                el = OxmlElement('w:'+side); el.set(qn('w:w'),val); el.set(qn('w:type'),'dxa'); margins.append(el)
            tcpr.append(margins)
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.05
            rich(par, str(text))
            if ci > 0 and re.fullmatch(r'[\d.,*—\s]+|n\. e\.', str(text)):
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in par.runs:
                run.font.size = Pt(size)
                if ri == 0: run.font.bold = True; run.font.color.rgb = RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def metric(key, scope):
    name = {'us':'volumeUS','de':'volumeDE','global':'volumeGlobal'}[scope]
    row = metrics[name].get(key)
    star = False
    if scope == 'us' and key in metrics['volumeUSFollow']:
        row = metrics['volumeUSFollow'][key]; star = True
    if row is None: return 'n. e.'
    value = row['avg_monthly_searches']
    text = '—' if value is None else f'{value:,}'.replace(',', '.')
    return text + ('*' if star else '')

def metric_table(keys, size=9):
    table(['Suchbegriff','USA EN','DE DE','Global EN'], [[k,metric(k,'us'),metric(k,'de'),metric(k,'global')] for k in keys], [.52,.16,.16,.16],size=size)

foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = foot.add_run('LinkedIn MCP Bridge · 30.10.2026 · '); r.font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); foot._p.append(field)

doc.add_paragraph(report['title'], 'Title')
doc.add_paragraph(report['subtitle'], 'Subtitle')
p(report['date'] + '  |  ' + report['product'] + '  |  Produktstand ' + report['baseline'])
p('Für Produktstrategie und Markteintritt. Die Studie verbindet direkte Google-Ads-Messungen mit öffentlichen Nutzerberichten und den vorhandenen Softwarefähigkeiten. Sie empfiehlt einen fokussierten Pilot für persönliche LinkedIn-Arbeit in lokalen AI-Agenten. Zahlungsbereitschaft und wiederholter Nutzen sind noch zu validieren.')
p('Datengrundlage: 7 frische MCP-Läufe · 1.053 Ergebniszeilen · 779 Suchbegriffe · 130 historisch geprüfte Begriffe · 31 öffentliche Quellen. Die Keywordwerte sind grobe Toolwerte und keine Marktgrößenschätzung.')
doc.add_heading('Leseweg', 2)
p('Entscheidung und Messung: Kapitel 1 bis 4. Zielgruppen und Aufgaben: Kapitel 5 bis 7. Produkt, Lücken und Wettbewerb: Kapitel 8 bis 11. Wirtschaftlichkeit, Prioritäten und Pilot: Kapitel 12 bis 16. Wiederverwendung und Entscheidung: Kapitel 17 und 18. Die Anhänge enthalten das Messprotokoll, alle historischen Suchbegriffe und das Quellenregister.')
for si, section in enumerate(report['sections']):
    heading = doc.add_heading(section['title'],1)
    if si == 0: heading.paragraph_format.page_break_before = True
    for block in section['blocks']:
        if block['type']=='p': p(block['text'])
        elif block['type']=='h2': doc.add_heading(block['text'],2)
        elif block['type']=='table': table(block['headers'],block['rows'],block.get('widths'))
        elif block['type']=='keyword_highlights':
            metric_table(['linkedin mcp','linkedin mcp server','claude linkedin connector','connect linkedin to claude','claude linkedin integration','connect linkedin to chatgpt','chatgpt linkedin integration','chatgpt linkedin connector','linkedin ai agent','linkedin inbox management','linkedin outreach automation','linkedin crm integration','linkedin post scheduler','linkedin scheduled posts','linkedin api','linkedin automation tools'])
            p('Gemeldete durchschnittliche Monatswerte, Oktober 2025 bis September 2026. * USA-Kontrolllauf vom selben Tag. — keine Metrik geliefert. n. e. nicht erhoben. Nicht addieren; die Scopes und Suchvarianten überlappen.')

heading = doc.add_heading('Anhang A Messprotokoll und Reproduzierbarkeit',1)
heading.paragraph_format.page_break_before = True
p('Die folgenden Dateien liegen neben dem Dokument im Ordner data. Sie enthalten die vollständigen normalisierten MCP-Antworten einschließlich roher Keywordfelder, Monatsreihen, Run-IDs, Scopes, Warnungen und ursprünglichen Artefaktpfaden. Die Konto-ID wurde redigiert. requests.json bewahrt die tatsächlich gesendeten Seeds und Scope-Anfragen; zur Wiederholung muss ein berechtigtes Ads-Konto verwendet werden.')
table(['Datei','Modus und Scope','Zeilen','Mit Wert'],[[n+'.json',('Historisch' if d['mode']=='historical_volume' else 'Discovery')+' · '+ ('Global EN' if d['resolved_scope']['locations'].get('all_locations') else ('DE DE' if d['resolved_scope']['language']['code']=='de' else 'USA EN')),str(len(d['results'])),str(sum(r['avg_monthly_searches'] is not None for r in d['results']))] for n,d in runs.items()],[.33,.35,.16,.16])
for n,d in runs.items():
    par=doc.add_paragraph(); par.paragraph_format.space_after=Pt(3)
    run=par.add_run(n+': ');run.bold=True
    run=par.add_run(d['run_id']);run.font.size=Pt(8.5)
p('Messfenster in den gelieferten Monatsreihen: 2025-10 bis 2026-09. Alle Läufe cache_hit=false. Original-Artefakte befinden sich zusätzlich im lokalen Browser-Bridge-Run-Store; die kopierten normalisierten Antworten sind der portable Datensatz. Die Metadatenbezeichnung raw_csv enthält in diesen Läufen JSON-Pfade und ist nicht als tatsächlicher CSV-Nachweis zu lesen.')
doc.add_heading('Kontrolllauf und fehlende Werte',2)
changes=[]
for k,b in metrics['volumeUSFollow'].items():
    a=metrics['volumeUS'].get(k)
    if a and a['avg_monthly_searches'] != b['avg_monthly_searches']:
        changes.append([k, '—' if a['avg_monthly_searches'] is None else str(a['avg_monthly_searches']), '—' if b['avg_monthly_searches'] is None else str(b['avg_monthly_searches'])])
table(['Suchbegriff','US zuerst','US Kontrolle'],changes,[.6,.2,.2])
p('Diese Differenzen begrenzen Aussagen aus einzelnen fehlenden Datenpunkten. Eine technische Ursache wurde nicht abschließend nachgewiesen. Die Studie verwendet im Vergleichsanhang den späteren historischen Wert, markiert ihn und erhält beide Rohantworten. Es wird kein Suchvolumen aus anderen Begriffen ergänzt.')
doc.add_heading('Repository Belege',2)
for text in ['README.md: Architektur, Installation, Funktionsgrenzen und Entwicklungsstatus.','src/core/contracts.ts: 27 Schemas, elf Mutationen, Eingaben und Limits.','src/core/capabilities.ts: Quellenadapter, Ergebnisverträge, Coverage und Acceptance-Zustände.','src/cli.ts: lokaler StdioServerTransport.','docs/validation.md: historische automatisierte und Live-Prüfungen; keine neue Acceptance durch diese Studie.','scripts/release-check.mjs und README.md: gesonderte Veröffentlichungsanforderungen.']:
    p(text)
p('Analysierter Produktcommit: 398095ebf0e8b2d82fd86720ab159ff908312811 auf lokalem main. Zum Recherchebeginn war kein Git-Remote konfiguriert und es gab keinen weiteren Worktree. Ein Remote-Fetch oder eine veröffentlichte Distribution ist damit nicht belegt. Der Bericht ändert keine Laufzeitfunktionen.')

heading = doc.add_heading('Anhang B Vollständiger historischer Keywordvergleich',1)
heading.paragraph_format.page_break_before = True
p('130 unterschiedliche historisch geprüfte Suchbegriffe, alphabetisch sortiert. Gemeldete Monatswerte für Oktober 2025 bis September 2026; grobe Stufen. Keine Summe bilden. USA EN, Deutschland DE und Global EN sind die angefragten Standort-/Sprachscopes. Global überlappt mit den Ländern und ist keine unabhängige Zusatznachfrage.')
p('* späterer USA-Kontrolllauf. — keine Metrik geliefert, kein Nachweis von null Nachfrage. n. e. in diesem Scope nicht erhoben. Die 779 insgesamt gefundenen Begriffe einschließlich Discovery stehen vollständig in den JSON-Daten.')
keys=sorted(set().union(*[set(metrics[n]) for n in ['volumeUS','volumeDE','volumeGlobal','volumeUSFollow']]))
metric_table(keys,size=8.5)

heading = doc.add_heading('Anhang C Quellenregister',1)
heading.paragraph_format.page_break_before = True
p('Alle Quellen wurden am 30. Oktober 2026 abgerufen. Die Kennungen im Text sind direkt verlinkt. Eine Produktbeschreibung belegt das Angebot des Anbieters, nicht seine unabhängige Wirksamkeit. Öffentliche Nutzerberichte sind qualitative Evidenz und keine repräsentativen Marktanteile.')
for s in sources:
    par=doc.add_paragraph();par.paragraph_format.keep_with_next=True
    par.paragraph_format.space_after=Pt(4)
    par.add_run(s['id']+'  ').bold=True
    hyperlink(par,s['title'],s['url'])
    par=doc.add_paragraph(s['type']+'. '+s['note']);par.paragraph_format.space_after=Pt(8)
    par.paragraph_format.line_spacing=1.05
    for run in par.runs: run.font.size=Pt(9)

out=ROOT/'linkedin-market-research.docx'
doc.save(out)
assert len(keys)==130
assert sum(len(d['results']) for d in runs.values())==1053
assert len(set(r['keyword'] for d in runs.values() for r in d['results']))==779
all_text=json.dumps(report,ensure_ascii=False)
assert set(re.findall(r'S\d{2}',all_text)) <= set(by_id)
print(json.dumps({'output':str(out),'sections':len(report['sections']),'sources':len(sources),'historical_keywords':len(keys),'control_differences':len(changes)},ensure_ascii=False))
